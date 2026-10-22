"""Read-only regression check for review R1; run from the superproject root."""
import io
import subprocess
from pathlib import Path

from openpyxl import load_workbook

baseline_commit = 'dfd422fa9211d4a2c40b740cb772b694779ada0f'
path = 'Team1_DTU/04_Quality_&_Testing/Test_Cases.xlsx'
source = subprocess.run(['git', '-C', 'docs', 'show', f'{baseline_commit}:{path}'], capture_output=True, check=True).stdout
baseline = load_workbook(io.BytesIO(source))
master = load_workbook(Path('docs') / path)
task = load_workbook('work/3.24-adherence-family-view/testcases.xlsx')

def values(sheet):
    return {cell.coordinate: cell.value for row in sheet for cell in row if cell.value is not None}

for sheet in baseline:
    if sheet.title == 'Report':
        continue
    assert sheet.title in master.sheetnames, sheet.title
    assert values(sheet) == values(master[sheet.title]), sheet.title
    assert str(sheet.merged_cells) == str(master[sheet.title].merged_cells), sheet.title
    print('PASS baseline cells/formulas/merges:', sheet.title)

assert values(task['Family adherence view']) == values(master['Family adherence view'])
assert {'Alert dispatch', 'Lịch sử tuân thủ', 'Family adherence view', 'Report'} <= set(master.sheetnames)
assert len(master.sheetnames) == 19
print('PASS: baseline 15 function sheets preserved; FAMADH matches task workbook; 19 total sheets')
